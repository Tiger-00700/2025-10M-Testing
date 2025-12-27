from docx import Document
from docx.oxml import parse_xml

doc = Document("1225全书定稿_final_standardized.docx")

print("Removing duplicate version lines...")
fixed_count = 0

for para in doc.paragraphs:
    text = para.text.strip()
    if text.startswith(">锚点名称:"):
        lines = text.split("\n")
        # Remove lines that are just "版本: v1.0" (without > prefix)
        filtered_lines = []
        for line in lines:
            if line.strip() == "版本: v1.0":
                continue  # Skip this duplicate line
            filtered_lines.append(line)
        
        if len(filtered_lines) != len(lines):
            # Lines were removed, update the paragraph
            new_text = "\n".join(filtered_lines)
            para.text = new_text
            fixed_count += 1

print(f"Removed duplicate version lines from {fixed_count} anchors")

# Re-apply visual formatting
formatted_count = 0
for para in doc.paragraphs:
    text = para.text.strip()
    if ("锚点名称:" in text) or ("锚点类型:" in text) or ("更新时间:" in text):
        p_element = para._element
        pPr = p_element.get_or_add_pPr()
        shading_xml = "<w:shd xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\" w:val=\"clear\" w:color=\"auto\" w:fill=\"E6F3FF\"/>"
        shading = parse_xml(shading_xml)
        pPr.append(shading)
        formatted_count += 1

doc.save("1225全书定稿_final_clean.docx")
print(f"Applied visual formatting to {formatted_count} anchor paragraphs")
print("Saved to: 1225全书定稿_final_clean.docx")
