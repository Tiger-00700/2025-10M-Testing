from docx import Document
import re

def renumber_anchors_in_docx(docx_path, output_path):
    """Renumber anchors in DOCX file by part, starting from 001 within each part"""
    doc = Document(docx_path)

    # First pass: collect all anchors with their positions and determine parts
    all_anchors = []

    current_part = 1  # Default to part 1

    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()

        # Check for part headers
        part_match = re.search(r'^第(\d+)篇', text)
        if part_match:
            current_part = int(part_match.group(1))

        # Check for anchor paragraphs
        if ('锚点类型:' in text or '锚点方法:' in text or
            text.startswith('> [锚点') or text.startswith('> [第') or
            text.startswith('[锚点') or text.startswith('[第')):

            # Extract anchor ID
            match = re.search(r'\[([^\]]+)\]', text)
            if match:
                anchor_id = match.group(1)

                all_anchors.append({
                    'para_idx': para_idx,
                    'original_id': anchor_id,
                    'part': current_part,
                    'text': text
                })

    # Group anchors by part
    anchors_by_part = {}
    for anchor in all_anchors:
        part = anchor['part']
        if part not in anchors_by_part:
            anchors_by_part[part] = []
        anchors_by_part[part].append(anchor)

    # Sort anchors within each part by their position
    for part in anchors_by_part:
        anchors_by_part[part].sort(key=lambda x: x['para_idx'])

    # Renumber anchors within each part
    total_renumbered = 0
    for part in sorted(anchors_by_part.keys()):
        anchors = anchors_by_part[part]
        for i, anchor_info in enumerate(anchors):
            new_anchor_id = f"第{part}篇-{(i+1):03d}"

            # Update the paragraph
            para = doc.paragraphs[anchor_info['para_idx']]
            old_text = anchor_info['text']
            new_text = old_text.replace(f'[{anchor_info["original_id"]}]', f'[{new_anchor_id}]')
            new_text = new_text.replace(f'**[{anchor_info["original_id"]}]**', f'**[{new_anchor_id}]**')
            para.text = new_text
            total_renumbered += 1

    doc.save(output_path)
    print(f"Successfully renumbered {total_renumbered} anchors in {len(anchors_by_part)} parts:")
    for part in sorted(anchors_by_part.keys()):
        count = len(anchors_by_part[part])
        print(f"  第{part}篇: {count} anchors (001-{count:03d})")

if __name__ == '__main__':
    renumber_anchors_in_docx('1225全书定稿_anchors_formatted.docx', '1225全书定稿_renumbered.docx')