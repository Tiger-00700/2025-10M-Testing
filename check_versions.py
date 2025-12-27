from docx import Document

doc = Document('1225全书定稿_final_standardized_fixed.docx')

# Check a few specific paragraphs to see the exact format
check_paragraphs = [15, 56, 120, 119]  # Include 119 which is our fixed anchor

for para_num in check_paragraphs:
    if para_num < len(doc.paragraphs):
        para = doc.paragraphs[para_num]
        text = para.text
        print(f'Paragraph {para_num}:')
        lines = text.split('\n')
        for i, line in enumerate(lines):
            if '版本' in line:
                print(f'  Line {i}: "{line}"')
        print('')