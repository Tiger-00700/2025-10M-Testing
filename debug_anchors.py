from docx import Document

doc = Document('1225全书定稿.docx')

print("Checking anchor formats...")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text.startswith('> [') and ']' in text:
        print(f"\nFound anchor at para {i}: {text}")

        # Check next few paragraphs
        for j in range(1, 8):
            if i + j < len(doc.paragraphs):
                next_text = doc.paragraphs[i + j].text.strip()
                print(f"  {j}: '{next_text}'")
                print(f"     starts with '> ': {next_text.startswith('> ')}")
                print(f"     contains ':': {':' in next_text}")
        break  # Just check first anchor