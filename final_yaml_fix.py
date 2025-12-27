from docx import Document

def fix_yaml_in_final_docx():
    doc = Document('1225全书定稿_final_complete.docx')

    # Find the specific YAML block and replace it
    for para_index, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if 'application_scenarios:' in text:
            # Replace this paragraph and the next few with proper YAML
            para.text = 'application_scenarios:'

            # Replace the next 16 paragraphs with proper YAML content
            yaml_content = [
                '  ecommerce_recommendation:',
                '    data_types: ["user_behavior_logs", "product_data"]',
                '    key_metrics: ["click_through_rate", "conversion_rate", "personalization_accuracy"]',
                '    script_reference: "examples/01_intro/simple_data_pipeline.md"',
                '  financial_risk_control:',
                '    data_types: ["transaction_records", "user_profiles"]',
                '    key_metrics: ["fraud_detection_rate", "false_positive_rate"]',
                '    script_reference: "examples/01_intro/data_factory.py"',
                '  iot_monitoring:',
                '    data_types: ["sensor_data", "device_status"]',
                '    key_metrics: ["data_integrity", "real_time_latency"]',
                '    script_reference: "examples/01_intro/masking_policy_demo.py"',
                '  content_distribution:',
                '    data_types: ["user_preferences", "content_metadata"]',
                '    key_metrics: ["content_matching_degree", "user_retention"]',
                '    script_reference: "examples/01_intro/qa_summary.py"'
            ]

            for i, yaml_line in enumerate(yaml_content):
                if para_index + i + 1 < len(doc.paragraphs):
                    doc.paragraphs[para_index + i + 1].text = yaml_line

            break

    doc.save('1225全书定稿_final_complete_fixed.docx')
    print('YAML content fixed in final DOCX')

if __name__ == '__main__':
    fix_yaml_in_final_docx()