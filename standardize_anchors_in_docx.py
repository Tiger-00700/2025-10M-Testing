from docx import Document
from docx.oxml import parse_xml
import re

def standardize_anchors_in_docx(input_path, output_path):
    """标准化原始DOCX文件中所有锚点的格式"""

    doc = Document(input_path)
    standardized_count = 0
    formatted_count = 0

    # 第一遍：收集所有锚点信息
    anchors = []

    for para in doc.paragraphs:
        text = para.text.strip()

        # 检查是否是锚点段落（支持三种格式）
        is_anchor = False
        anchor_id = ""
        title = ""

        # 格式1: > [锚点XXX] 标题 > 锚点类型: > 锚点方法:
        if text.startswith('> [') and '> 锚点类型:' in text and '> 锚点方法:' in text:
            is_anchor = True
            id_match = re.search(r'> \[([^\]]+)\]', text)
            if id_match:
                anchor_id = id_match.group(1)
            title_match = re.search(r'\] ([^>]+)', text)
            title = title_match.group(1).strip() if title_match else ""

        # 格式2: [锚点XXX] 标题 锚点类型: 锚点方法:
        elif text.startswith('[锚点') and '锚点类型:' in text and '锚点方法:' in text and not text.startswith('>'):
            is_anchor = True
            id_match = re.search(r'\[([^\]]+)\]', text)
            if id_match:
                anchor_id = id_match.group(1)
            title_match = re.search(r'\] ([^\n]+)', text)
            title = title_match.group(1).strip() if title_match else ""

        # 格式3: [第X篇-XXX] 标题 锚点类型: 锚点方法:
        elif re.match(r'^\[第\d+篇-\d+\]', text) and '锚点类型:' in text and '锚点方法:' in text:
            is_anchor = True
            id_match = re.search(r'\[([^\]]+)\]', text)
            if id_match:
                anchor_id = id_match.group(1)
            title_match = re.search(r'\] ([^\n]+)', text)
            title = title_match.group(1).strip() if title_match else ""

        # 特殊格式: [第X篇-XXX] 标题 锚点类型: ... (没有锚点方法:)
        elif re.match(r'^\[第\d+篇-\d+\]', text) and '锚点类型:' in text and '锚点方法:' not in text:
            is_anchor = True
            id_match = re.search(r'\[([^\]]+)\]', text)
            if id_match:
                anchor_id = id_match.group(1)
            title_match = re.search(r'\] ([^锚点]+)', text)
            title = title_match.group(1).strip() if title_match else ""

        # 特殊格式: [锚点XX-XXX] 标题 锚点类型: ... 锚点名称: ... (一行包含所有字段)
        elif re.match(r'^\[锚点\d+-\d+\]', text) and '锚点类型:' in text and '锚点方法:' in text and '锚点名称:' in text:
            is_anchor = True
            id_match = re.search(r'\[([^\]]+)\]', text)
            if id_match:
                anchor_id = id_match.group(1)
            # 对于这种特殊格式，从第二个"锚点名称:"后提取标题
            name_match = re.findall(r'锚点名称:\s*([^锚点]+?)(?=\s*锚点类型:|$)', text)
            if len(name_match) >= 2:
                title = name_match[1].strip()
            else:
                title_match = re.search(r'\] ([^锚点]+)', text)
                title = title_match.group(1).strip() if title_match else ""

        if is_anchor:
            print(f"Found anchor: {anchor_id} - {title[:30]}...")
            anchors.append({
                'id': anchor_id,
                'title': title,
                'paragraph': para,
                'original_text': text
            })

    print(f"Found {len(anchors)} anchors to process")

    # 第二遍：标准化锚点格式
    for anchor in anchors:
        print(f"Processing anchor: {anchor['id']}")

        # 从原始文本中提取字段
        original_text = anchor['original_text']

        fields = {}
        # 提取各个字段（支持有>前缀和无>前缀的格式，以及特殊格式）
        # 对于特殊格式，字段可能在"锚点名称:"之后
        if '锚点名称:' in original_text and '锚点类型:' in original_text:
            # 特殊格式：提取第一个锚点名称后的字段
            parts = original_text.split('锚点名称:')
            if len(parts) >= 2:
                # 取第一个锚点名称之后的内容
                field_text = '锚点名称:' + parts[1]
            else:
                field_text = original_text
        else:
            field_text = original_text

        # 解析各个字段 - 使用更精确的正则表达式
        type_match = re.search(r'锚点类型:\s*([^锚点]+?)(?=锚点方法:|引用附件:|完成情况:|补充建议:|书稿版本:|更新时间:|$)', field_text)
        if type_match:
            fields['type'] = type_match.group(1).strip()

        method_match = re.search(r'锚点方法:\s*([^锚点]+?)(?=引用附件:|完成情况:|补充建议:|书稿版本:|更新时间:|$)', field_text)
        if method_match:
            fields['method'] = method_match.group(1).strip()

        ref_match = re.search(r'引用附件:\s*([^锚点]+?)(?=完成情况:|补充建议:|书稿版本:|更新时间:|$)', field_text)
        if ref_match:
            fields['reference'] = ref_match.group(1).strip()

        status_match = re.search(r'完成情况:\s*([^锚点]+?)(?=补充建议:|书稿版本:|更新时间:|$)', field_text)
        if status_match:
            fields['status'] = status_match.group(1).strip()

        suggestion_match = re.search(r'补充建议:\s*([^锚点]+?)(?=引用附件:|完成情况:|书稿版本:|更新时间:|$)', field_text)
        if suggestion_match:
            fields['suggestion'] = suggestion_match.group(1).strip()

        version_match = re.search(r'书稿版本:\s*([^锚点]+?)(?=更新时间:|$)', field_text)
        if version_match:
            fields['version'] = version_match.group(1).strip()

        time_match = re.search(r'更新时间:\s*([^锚点]+)', field_text)
        if time_match:
            fields['timestamp'] = time_match.group(1).strip()

        # 确定篇章编号
        chapter_prefix = ""
        if '第1篇' in anchor['id'] or anchor['id'].startswith('锚点'):
            chapter_prefix = "第1篇"
        elif '第2篇' in anchor['id']:
            chapter_prefix = "第2篇"
        elif '第3篇' in anchor['id']:
            chapter_prefix = "第3篇"
        elif '第4篇' in anchor['id']:
            chapter_prefix = "第4篇"
        elif '第5篇' in anchor['id']:
            chapter_prefix = "第5篇"
        elif '第6篇' in anchor['id']:
            chapter_prefix = "第6篇"
        elif '第7篇' in anchor['id']:
            chapter_prefix = "第7篇"

        # 生成标准化的锚点文本（新格式）
        # 清理ID中的篇章前缀，用于生成新的标准格式
        clean_id = anchor['id']
        clean_id = re.sub(r'^第\d+篇-', '', clean_id)
        clean_id = re.sub(r'^锚点', '', clean_id)

        anchor_name = f"{chapter_prefix}[{clean_id}]-{anchor['title']}"

        # 构建标准化文本行
        standardized_lines = [
            f'>锚点名称: {anchor_name}',
            f'>锚点类型: {fields.get("type", "建议")}',
            f'>锚点方法: {fields.get("method", "配置示例")}',
            f'>补充建议: {fields.get("suggestion", "")}',
            f'>引用附件: {fields.get("reference", "")}',
            f'>完成情况: {fields.get("status", "已标准化")}',
        ]
        
        # 只有在原始内容中没有版本信息时才添加版本字段
        if not fields.get("version"):
            standardized_lines.append(f'>书稿版本: v1.0')
        
        standardized_lines.append(f'>更新时间: {fields.get("timestamp", "2025-12-26")}')

        standardized_text = '\n'.join(standardized_lines)

        # 完全重写段落文本，清除所有旧内容
        anchor['paragraph'].text = standardized_text
        print(f"Paragraph updated - new text length: {len(standardized_text)}")
        standardized_count += 1

    # 应用视觉格式到所有锚点段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if ('锚点名称:' in text) or ('锚点类型:' in text) or ('更新时间:' in text):
            # 应用浅蓝色背景
            p_element = para._element
            pPr = p_element.get_or_add_pPr()
            shading_xml = '''<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="E6F3FF"/>'''
            shading = parse_xml(shading_xml)
            pPr.append(shading)
            formatted_count += 1

    doc.save(output_path)
    print(f"Standardized {standardized_count} anchor blocks")
    print(f"Applied visual formatting to {formatted_count} anchor paragraphs")
    print(f"Saved to: {output_path}")

if __name__ == '__main__':
    standardize_anchors_in_docx('1225全书定稿.docx', '1225全书定稿_final_standardized.docx')